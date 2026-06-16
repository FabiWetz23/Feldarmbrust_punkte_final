from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.drawing.image import Image as ExcelImage
from io import BytesIO
from pathlib import Path
from typing import List, Dict, Any
from datetime import datetime
from .models import Competition, Series, Shooter


class ExcelExporter:
    def __init__(self, competition: Competition, sponsor_paths: list[Path] | None = None, ko_state: dict | None = None):
        self.competition = competition
        self.sponsor_paths = sponsor_paths or []
        self.ko_state = ko_state or {}
        self.wb = Workbook()
        self.ws = self.wb.active
        self.ws.title = "Competition Results"
        self.header_row = 1
        self.data_start_row = 2
        
    def create_styling(self):
        """Erstellt die Formatvorlagen"""
        # Überschrift Format
        self.header_font = Font(name='Arial', size=12, bold=True, color='FFFFFF')
        self.header_fill = PatternFill(start_color='366092', end_color='366092', fill_type='solid')
        self.header_alignment = Alignment(horizontal='center', vertical='center')
        
        # Daten Format
        self.data_font = Font(name='Arial', size=10)
        self.data_alignment = Alignment(horizontal='center', vertical='center')
        
        # Rahmen
        self.thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        # Summen Format (fett)
        self.sum_font = Font(name='Arial', size=10, bold=True)
        self.sum_fill = PatternFill(start_color='F2F2F2', end_color='F2F2F2', fill_type='solid')
    def setup_headers(self, max_shots: int):
        """Creates the table headers"""
        headers = ['Rank', 'No', 'Competitor', 'Country', 'Round (End)']
        
        # Shot columns
        for i in range(1, max_shots + 1):
            headers.append(f'Shot {i}')
        
        headers.extend(['End Sum', 'Total Sum', 'Inner Tens (X)'])
        
        # Überschriften einfügen
        for col, header in enumerate(headers, 1):
            cell = self.ws.cell(row=self.header_row, column=col, value=header)
            cell.font = self.header_font
            cell.fill = self.header_fill
            cell.alignment = self.header_alignment
            cell.border = self.thin_border
    
    def fill_data(self):
        """Füllt die Tabelle mit Daten"""
        row = self.data_start_row
        
        # Gruppiere Serien nach Schützen und Durchgängen
        shooter_data = {}
        max_shots = 0
        
        for series in self.competition.series.values():
            shooter_id = series.shooter_id
            if shooter_id not in shooter_data:
                shooter_data[shooter_id] = {}
            
            shooter_data[shooter_id][series.round_number] = series
            max_shots = max(max_shots, series.shots_per_series)
        
        # Daten für jeden Schützen einfügen
        # Create a sorted leaderboard to get the rank
        leaderboard = []
        for sid, sh in self.competition.shooters.items():
            total_score = sum(
                sum(10.0 if shot.score == 11.0 else shot.score for shot in s.shots)
                for s in self.competition.series.values()
                if s.shooter_id == sid and s.round_number >= 1
            )
            inner_tens = sum(
                1 for s in self.competition.series.values()
                if s.shooter_id == sid
                for shot in s.shots
                if shot.score == 11.0
            )
            leaderboard.append({
                "id": sid,
                "name": sh.name,
                "total_score": total_score,
                "inner_tens": inner_tens
            })
        
        # Sort: total desc, then inner tens desc, then name asc
        leaderboard.sort(key=lambda x: (-x["total_score"], -x["inner_tens"], x["name"]))
        
        shooter_ranks = {entry["id"]: i + 1 for i, entry in enumerate(leaderboard)}

        # Daten für jeden Schützen einfügen (sortiert nach Rank)
        for entry in leaderboard:
            shooter_id = entry["id"]
            shooter = self.competition.shooters[shooter_id]
            rank = shooter_ranks.get(shooter_id, "-")
            
            if shooter_id not in shooter_data:
                continue
                
            series_dict = shooter_data[shooter_id]
            total_sum = 0
            
            # Für jeden Durchgang eine Zeile (sortiert nach Round Number)
            for round_num in sorted(series_dict.keys()):
                series = series_dict[round_num]
                
                # Calculate series sum (11 counts as 10)
                series_sum = sum(10.0 if shot.score == 11.0 else shot.score for shot in series.shots)
                
                # Only add to total for match rounds (>= 1)
                if round_num >= 1:
                    total_sum += series_sum
                
                # Schützen-Info (nur in erster Zeile)
                rank_cell = self.ws.cell(row=row, column=1, value=rank)
                no_cell = self.ws.cell(row=row, column=2, value=getattr(shooter, "start_number", ""))
                name_cell = self.ws.cell(row=row, column=3, value=shooter.name)
                country_cell = self.ws.cell(row=row, column=4, value=getattr(shooter, "country", ""))
                
                if round_num == min(series_dict.keys()):
                    rank_cell.font = self.data_font
                    no_cell.font = self.data_font
                    name_cell.font = self.data_font
                    country_cell.font = self.data_font
                else:
                    rank_cell.value = ""
                    no_cell.value = ""
                    name_cell.value = ""
                    country_cell.value = ""
                
                # Durchgang mit Typ
                if round_num == -1:
                    round_display = "Sighting 1"
                elif round_num == 0:
                    round_display = "Sighting 2"
                else:
                    round_display = f"End {round_num}"
                
                self.ws.cell(row=row, column=5, value=round_display)
                
                # Schüsse mit X für Inner Tens (Score 11)
                for shot in series.shots:
                    col = 5 + shot.shot_number
                    display_score = "X" if shot.score == 11.0 else str(shot.score)
                    self.ws.cell(row=row, column=col, value=display_score)
                
                # Serien-Summe
                sum_col = 5 + max_shots + 1
                self.ws.cell(row=row, column=sum_col, value=series_sum)
                
                # Gesamt-Summe (nur in letzter Zeile des Schützen für Match Rounds)
                if round_num == max([r for r in series_dict.keys() if r >= 1]):
                    total_col = 5 + max_shots + 2
                    self.ws.cell(row=row, column=total_col, value=total_sum)
                    
                    x_col = 5 + max_shots + 3
                    inner_tens = entry.get("inner_tens", 0)
                    self.ws.cell(row=row, column=x_col, value=inner_tens)
                
                # Formatierung der Datenzeile
                self._format_data_row(row, max_shots, round_num == max([r for r in series_dict.keys() if r >= 1]))
                row += 1
            
            # Leere Zeile zwischen Schützen
            row += 1
    
    def _format_data_row(self, row: int, max_shots: int, is_last_round: bool):
        """Formatiert eine Datenzeile"""
        # Alle Zellen in der Zeile formatieren (inkl. Inner Tens Spalte)
        for col in range(1, 5 + max_shots + 4):
            cell = self.ws.cell(row=row, column=col)
            cell.font = self.data_font
            cell.alignment = self.data_alignment
            cell.border = self.thin_border
            
            # Summen-Spalten hervorheben
            if col in [5 + max_shots + 1, 5 + max_shots + 2, 5 + max_shots + 3]:
                if is_last_round or col == 5 + max_shots + 1:  # Serien-Summe immer, Gesamt-Summe & X nur letzte Zeile
                    cell.font = self.sum_font
                    cell.fill = self.sum_fill
    
    def adjust_column_widths(self, max_shots: int):
        """Passt die Spaltenbreiten an"""
        # Rank, No, Name und Verein
        self.ws.column_dimensions['A'].width = 8
        self.ws.column_dimensions['B'].width = 8
        self.ws.column_dimensions['C'].width = 20
        self.ws.column_dimensions['D'].width = 25
        
        # Durchgang
        self.ws.column_dimensions['E'].width = 15
        
        # Schuss-Spalten
        for i in range(1, max_shots + 1):
            col_letter = get_column_letter(5 + i)
            self.ws.column_dimensions[col_letter].width = 8
            
        # Summen und X
        self.ws.column_dimensions[get_column_letter(5 + max_shots + 1)].width = 12
        self.ws.column_dimensions[get_column_letter(5 + max_shots + 2)].width = 12
        self.ws.column_dimensions[get_column_letter(5 + max_shots + 3)].width = 12
        
        # Summen-Spalten
        sum_col_letter = get_column_letter(5 + max_shots + 1)
        total_col_letter = get_column_letter(5 + max_shots + 2)
        self.ws.column_dimensions[sum_col_letter].width = 12
        self.ws.column_dimensions[total_col_letter].width = 12
    
    def add_metadata(self):
        """Fügt Metadaten hinzu"""
        max_cols = 5 + self._get_max_shots() + 3
        if self.sponsor_paths:
            self._add_sponsors(max_cols)
            title_row = 3
            date_row = 4
            self.header_row = 6
        else:
            title_row = 1
            date_row = 2
            self.header_row = 4
        self.data_start_row = self.header_row + 1
        
        # Titel über der Tabelle
        title_cell = self.ws.cell(row=title_row, column=1)
        title_cell.value = self.competition.name
        title_cell.font = Font(name='Arial', size=16, bold=True)
        title_cell.alignment = Alignment(horizontal='center', vertical='center')
        self.ws.merge_cells(start_row=title_row, start_column=1, end_row=title_row, end_column=max_cols)
        
        # Datum und Uhrzeit
        date_cell = self.ws.cell(row=date_row, column=1)
        date_cell.value = f"Exportiert: {datetime.now().strftime('%d.%m.%Y %H:%M:%S')}"
        date_cell.font = Font(name='Arial', size=10, italic=True)
        date_cell.alignment = Alignment(horizontal='center', vertical='center')
        self.ws.merge_cells(start_row=date_row, start_column=1, end_row=date_row, end_column=max_cols)

    def _add_sponsors(self, max_cols: int):
        """Adds sponsor logos to the top of the worksheet."""
        self.ws.row_dimensions[1].height = 58
        label = self.ws.cell(row=1, column=1, value="Sponsoren")
        label.font = Font(name='Arial', size=10, bold=True)
        label.alignment = Alignment(horizontal='left', vertical='center')

        col = 2
        for path in self.sponsor_paths[:8]:
            try:
                img = ExcelImage(str(path))
                if img.width > 180:
                    ratio = 180 / img.width
                    img.width = 180
                    img.height = max(1, int(img.height * ratio))
                if img.height > 62:
                    ratio = 62 / img.height
                    img.height = 62
                    img.width = max(1, int(img.width * ratio))
                self.ws.add_image(img, f"{get_column_letter(col)}1")
                col += 2
                if col > max_cols:
                    break
            except Exception:
                continue
    
    def _get_max_shots(self) -> int:
        """Ermittelt die maximale Anzahl an Schüssen pro Serie"""
        max_shots = 0
        for series in self.competition.series.values():
            max_shots = max(max_shots, series.shots_per_series)
        return max_shots

    @staticmethod
    def _ko_score_text(value) -> str:
        if value is None:
            return ""
        return "X" if int(value) == 11 else str(value)

    @staticmethod
    def _ko_score_points(value) -> int:
        if value is None:
            return 0
        value = int(value)
        return 10 if value == 11 else value

    def _ko_overall_winner(self):
        rounds = self.ko_state.get("rounds") or []
        if not rounds:
            return None
        final_match = (rounds[-1].get("matches") or [None])[0]
        if not final_match:
            return None
        winner_slot = final_match.get("winner_slot")
        if winner_slot is None:
            return None
        competitors = final_match.get("competitors") or []
        if winner_slot >= len(competitors):
            return None
        return competitors[winner_slot]

    def _ko_end_total(self, ends: list, end_index: int) -> str:
        if end_index >= len(ends):
            return ""
        shots = ends[end_index] or []
        if not any(value is not None for value in shots):
            return ""
        total = sum(self._ko_score_points(value) for value in shots)
        inner_tens = sum(1 for value in shots if value == 11)
        return f"{total}" + (f"/{inner_tens}X" if inner_tens else "")

    def _set_cell_border(self, ws, row: int, col: int, **sides) -> None:
        cell = ws.cell(row=row, column=col)
        cell.border = Border(
            left=sides.get("left", cell.border.left),
            right=sides.get("right", cell.border.right),
            top=sides.get("top", cell.border.top),
            bottom=sides.get("bottom", cell.border.bottom),
        )

    def _draw_ko_connector(self, ws, source_row: int, target_row: int, start_col: int, end_col: int) -> None:
        line = Side(style="medium", color="404040")
        if end_col <= start_col:
            return

        for col in range(start_col, end_col + 1):
            self._set_cell_border(ws, source_row, col, top=line)
        for row in range(min(source_row, target_row), max(source_row, target_row) + 1):
            self._set_cell_border(ws, row, end_col, right=line)
        for col in range(end_col, end_col + 3):
            self._set_cell_border(ws, target_row, col, top=line)

    def _draw_ko_match_box(self, ws, row: int, col: int, round_name: str, match: dict) -> None:
        dark_fill = PatternFill(start_color="262626", end_color="262626", fill_type="solid")
        win_fill = PatternFill(start_color="E2F0D9", end_color="E2F0D9", fill_type="solid")
        no_fill = PatternFill(start_color="FFF176", end_color="FFF176", fill_type="solid")
        light_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
        header_fill = PatternFill(start_color="D9EAF7", end_color="D9EAF7", fill_type="solid")
        total_fill = PatternFill(start_color="595959", end_color="595959", fill_type="solid")
        white_font = Font(name="Arial", size=8, bold=True, color="FFFFFF")
        name_font = Font(name="Arial", size=8, bold=True)
        tiny_font = Font(name="Arial", size=7)

        box_cols = 13
        headers = ["", "No", "Competitor", "", "", "", "E1", "E2", "E3", "E4", "E5", "SO", "Pts"]
        for offset, label in enumerate(headers):
            cell = ws.cell(row=row - 1, column=col + offset, value=label)
            cell.fill = header_fill
            cell.font = tiny_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = self.thin_border

        competitors = match.get("competitors") or [None, None]
        points = match.get("match_points") or [0, 0]
        winner_slot = match.get("winner_slot")

        for slot in (0, 1):
            r = row + slot
            competitor = competitors[slot] if slot < len(competitors) else None
            ends = (match.get("ends") or [[], []])[slot] if slot < len(match.get("ends") or []) else []
            is_winner = winner_slot == slot

            for c in range(col, col + box_cols):
                cell = ws.cell(row=r, column=c)
                cell.fill = win_fill if is_winner else light_fill
                cell.border = self.thin_border
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                cell.font = tiny_font

            ws.cell(row=r, column=col, value=match.get("match_number") if slot == 0 else "")
            ws.cell(row=r, column=col).fill = dark_fill
            ws.cell(row=r, column=col).font = white_font

            ws.cell(row=r, column=col + 1, value=competitor.get("start_number") if competitor else "")
            ws.cell(row=r, column=col + 1).fill = no_fill if competitor else light_fill
            ws.cell(row=r, column=col + 1).font = name_font

            name = competitor.get("name") if competitor else "Free"
            ws.merge_cells(start_row=r, start_column=col + 2, end_row=r, end_column=col + 5)
            name_cell = ws.cell(row=r, column=col + 2, value=name)
            name_cell.font = name_font
            name_cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

            for end_index in range(5):
                ws.cell(row=r, column=col + 6 + end_index, value=self._ko_end_total(ends, end_index))

            shoot_off = (match.get("shoot_off") or [None, None])[slot]
            ws.cell(row=r, column=col + 11, value=self._ko_score_text(shoot_off))
            pts = ws.cell(row=r, column=col + 12, value=points[slot] if slot < len(points) else 0)
            pts.fill = total_fill
            pts.font = white_font

        ws.cell(row=row + 2, column=col, value=round_name).font = Font(name="Arial", size=7, italic=True, color="666666")
        ws.merge_cells(start_row=row + 2, start_column=col, end_row=row + 2, end_column=col + box_cols - 1)

    def add_ko_sheet(self):
        ws = self.wb.create_sheet("KO Bracket")
        ws.cell(row=1, column=1, value="KO Bracket").font = Font(name="Arial", size=16, bold=True)
        if not self.ko_state.get("rounds"):
            ws.cell(row=2, column=1, value="No KO bracket available.").font = Font(name="Arial", size=12, italic=True)
            ws.column_dimensions["A"].width = 32
            return

        winner = self._ko_overall_winner()
        ws.cell(row=2, column=1, value=self.competition.name).font = Font(name="Arial", size=12, bold=True)
        ws.cell(
            row=3,
            column=1,
            value=f"Overall KO Winner: {winner.get('start_number') or '-'} {winner.get('name')}" if winner else "Overall KO Winner: -",
        ).font = Font(name="Arial", size=11, bold=True)

        rounds = self.ko_state.get("rounds") or []
        box_width = 13
        col_step = 15
        row_gap = 5
        start_row = 7
        start_col = 1
        row_positions = []

        first_count = max(1, len(rounds[0].get("matches", [])))
        row_positions.append([start_row + i * row_gap for i in range(first_count)])
        for round_index in range(1, len(rounds)):
            positions = []
            for match_index, _match in enumerate(rounds[round_index].get("matches", [])):
                left = row_positions[round_index - 1][match_index * 2]
                right_index = min(match_index * 2 + 1, len(row_positions[round_index - 1]) - 1)
                right = row_positions[round_index - 1][right_index]
                positions.append((left + right) // 2)
            row_positions.append(positions)

        max_col = start_col + (len(rounds) - 1) * col_step + box_width
        max_row = max(max(rows) for rows in row_positions if rows) + 4
        for col in range(1, max_col + 3):
            ws.column_dimensions[get_column_letter(col)].width = 3.8
        for round_index in range(len(rounds)):
            col = start_col + round_index * col_step
            ws.column_dimensions[get_column_letter(col + 1)].width = 5.5
            for name_col in range(col + 2, col + 6):
                ws.column_dimensions[get_column_letter(name_col)].width = 8.5
            for score_col in range(col + 6, col + box_width):
                ws.column_dimensions[get_column_letter(score_col)].width = 4.8
        for row in range(1, max_row + 2):
            ws.row_dimensions[row].height = 18

        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max_col)
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=max_col)
        ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=max_col)
        ws.cell(row=1, column=1).alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(row=2, column=1).alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(row=3, column=1).alignment = Alignment(horizontal="center", vertical="center")

        for round_index, round_data in enumerate(rounds):
            col = start_col + round_index * col_step
            title = ws.cell(row=5, column=col, value=round_data.get("name"))
            title.font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
            title.fill = self.header_fill
            title.alignment = Alignment(horizontal="center", vertical="center")
            ws.merge_cells(start_row=5, start_column=col, end_row=5, end_column=col + box_width - 1)

            for match_index, match in enumerate(round_data.get("matches", [])):
                row = row_positions[round_index][match_index]
                self._draw_ko_match_box(ws, row, col, round_data.get("name"), match)

        for round_index in range(len(rounds) - 1):
            source_col = start_col + round_index * col_step
            connector_start = source_col + box_width
            connector_end = source_col + col_step - 1
            for source_index, source_row in enumerate(row_positions[round_index]):
                target_index = source_index // 2
                if target_index < len(row_positions[round_index + 1]):
                    target_row = row_positions[round_index + 1][target_index]
                    self._draw_ko_connector(ws, source_row, target_row, connector_start, connector_end)

        legend_row = max_row + 1
        ws.cell(row=legend_row, column=1, value="Legend: E1-E5 = end totals, /X = inner tens, SO = shoot-off, Pts = match points.").font = Font(name="Arial", size=9, italic=True)
        ws.merge_cells(start_row=legend_row, start_column=1, end_row=legend_row, end_column=min(max_col, 18))
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 1
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.sheet_view.showGridLines = False
    
    def generate_excel(self) -> BytesIO:
        """Generiert die Excel-Datei"""
        self.create_styling()
        max_shots = self._get_max_shots()
        
        self.add_metadata()
        self.setup_headers(max_shots)
        self.fill_data()
        self.adjust_column_widths(max_shots)
        self.add_ko_sheet()
        
        # In BytesIO speichern
        excel_buffer = BytesIO()
        self.wb.save(excel_buffer)
        excel_buffer.seek(0)
        
        return excel_buffer


def export_competition_to_excel(
    competition: Competition,
    sponsor_paths: list[Path] | None = None,
    ko_state: dict | None = None,
) -> BytesIO:
    """Exportiert einen Wettkampf nach Excel"""
    exporter = ExcelExporter(competition, sponsor_paths, ko_state)
    return exporter.generate_excel()
